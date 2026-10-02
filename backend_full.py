import os
import json
import io
import pandas as pd
import docx
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from Bio import Entrez
from langchain_openai import ChatOpenAI
from langchain_core.messages import SystemMessage, HumanMessage
from dotenv import load_dotenv

load_dotenv()

app = FastAPI(title="Apex Athlete Multi-Agent Hub")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

Entrez.email = os.getenv("NCBI_EMAIL", "sporcu@universite.edu.tr")
llm = ChatOpenAI(model="gpt-4o", temperature=0.1)

# --- SISTEM PROMPTLARI ---

KOC_PROMPT = """
Sen otonom bir spor fizyoloğu ve beslenme koçusun.
KATI KURAL: Sadece antrenman, toparlanma ve sporcu beslenmesi konuşursun.
Çıktıyı SADECE şu JSON şemasında ver:
{
  "toparlanma_skoru": 0-100,
  "durum": "Mükemmel / Dikkat / Tükenmiş",
  "biyolojik_mekanizma": "Sade, metaforlu biyolojik süreç",
  "eski_ekol_miti": "1990-2000'lerin eski yanlış kabulü",
  "guncel_bilim": "Son 5-10 yılın modern bilimsel gerçeği",
  "koc_direktifi": "Bugün yapılması gereken net eylem",
  "beslenme_toparlanma": "Makro, kalori ve mikro besin planı",
  "kaynakca": ["Yazar, Yıl, Başlık, DOI/PMID"]
}
"""

FORM_PROMPT = """
Sen otonom bir biyomekanik ve hareket formu uzmanısın.
Kullanıcının hatasını analiz et ve hatasız örnek formu inşa et.
Çıktıyı SADECE şu JSON şemasında ver:
{
  "hareket": "Egzersiz adı",
  "hata_tespiti": "Kullanıcının yaptığı biyomekanik sapma",
  "sakatlik_riski": "Yük binen tendon/eklem/bağ doku",
  "kusursuz_ornek_form": {
    "ayak_zemin": "Ayak açısı ve zemin torku",
    "diz_kalca": "Diz takip hattı ve kalça derinliği",
    "govde_omurga": "Omurga açısı ve karın içi basınç",
    "bar_yolu": "Dikey eksen ve mid-foot hizası"
  },
  "zihinsel_komut": "Sete girerken akılda tutulacak 2 kelimelik cue"
}
"""

PROGRAM_PROMPT = """
Sen kanıta dayalı (evidence-based) antrenman programı denetleyicisisin.
Görevin yüklenen Excel/Word antrenman programını okumak, hacim (set/hafta), frekans, RPE/RIR dağılımı ve egzersiz seçimini analiz etmektir.
Çıktıyı SADECE şu JSON şemasında ver:
{
  "program_ozeti": "Programın genel yapısı ve haftalık frekansı",
  "haftalik_hacim_analizi": "Kas grubu başına düşen set sayısı ve değerlendirmesi",
  "eski_vs_yeni_yorumu": "Bu programdaki eski ekol alışkanlıklar ve modern bilimin önerdiği düzeltmeler",
  "asiri_veya_eksik_noktalar": "Junk volume (çöp hacim), toparlanma riski veya yetersiz uyarılan bölgeler",
  "kocun_revizyon_onerisi": "Programa yapılacak doğrudan cerrahi müdahale",
  "bilimsel_dayanak": ["Schoenfeld, Israetel vb. hacim ve hipertrofi çalışmaları"]
}
"""

def pubmed_ara(query: str):
    try:
        q = f"{query} AND (resistance training OR muscle hypertrophy OR exercise physiology)"
        h = Entrez.esearch(db="pubmed", term=q, retmax=2, sort="relevance")
        res = Entrez.read(h)
        h.close()
        ids = res.get("IdList", [])
        if not ids:
            return "Literatür özeti bulunamadı."
        fh = Entrez.efetch(db="pubmed", id=",".join(ids), rettype="medline", retmode="text")
        data = fh.read()
        fh.close()
        return data[:2000]
    except Exception:
        return "PubMed araması yapılamadı."

# --- ENDPOINT TANIMLARI ---

class KocTalep(BaseModel):
    soru: str
    rpe: float = 8.0
    uyku: float = 7.0

class FormTalep(BaseModel):
    hareket: str
    hata_sikayet: str

@app.post("/ajan/koc")
async def ajan_koc(talep: KocTalep):
    pubmed_veri = pubmed_ara(talep.soru)
    prompt = f"Soru: {talep.soru}\nRPE: {talep.rpe}\nUyku: {talep.uyku} saat\nPubMed Verisi:\n{pubmed_veri}"
    res = await llm.ainvoke([SystemMessage(content=KOC_PROMPT), HumanMessage(content=prompt)])
    clean = res.content.replace("```json", "").replace("```", "").strip()
    return json.loads(clean)

@app.post("/ajan/form")
async def ajan_form(talep: FormTalep):
    prompt = f"Hareket: {talep.hareket}\nHata/Ağrı Bildirimi: {talep.hata_sikayet}"
    res = await llm.ainvoke([SystemMessage(content=FORM_PROMPT), HumanMessage(content=prompt)])
    clean = res.content.replace("```json", "").replace("```", "").strip()
    return json.loads(clean)

@app.post("/ajan/program-incele")
async def ajan_program_incele(dosya: UploadFile = File(...)):
    icerik_metni = ""
    dosya_adi = dosya.filename.lower()
    dosya_bayt = await dosya.read()

    # Excel Okuma (.xlsx, .xls)
    if dosya_adi.endswith(('.xlsx', '.xls')):
        try:
            excel_verisi = pd.read_excel(io.BytesIO(dosya_bayt), sheet_name=None)
            for sayfa, df in excel_verisi.items():
                icerik_metni += f"\n--- Sayfa: {sayfa} ---\n"
                icerik_metni += df.dropna(how='all').to_string()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Excel okuma hatası: {str(e)}")

    # Word Okuma (.docx)
    elif dosya_adi.endswith('.docx'):
        try:
            doc = docx.Document(io.BytesIO(dosya_bayt))
            paragraflar = [p.text for p in doc.paragraphs if p.text.strip()]
            tablolar = []
            for tablo in doc.tables:
                for row in tablo.rows:
                    tablolar.append(" | ".join([cell.text.strip() for cell in row.cells]))
            icerik_metni = "\n".join(paragraflar) + "\n\nTABLOLAR:\n" + "\n".join(tablolar)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Word okuma hatası: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Sadece .xlsx, .xls veya .docx formatında dosya yükleyebilirsiniz.")

    if not icerik_metni.strip():
        raise HTTPException(status_code=400, detail="Dosya içeriği boş veya okunamadı.")

    res = await llm.ainvoke([
        SystemMessage(content=PROGRAM_PROMPT),
        HumanMessage(content=f"Aşağıdaki antrenman programını modern spor bilimi ışığında denetle:\n\n{icerik_metni[:10000]}")
    ])
    clean = res.content.replace("```json", "").replace("```", "").strip()
    return json.loads(clean)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
